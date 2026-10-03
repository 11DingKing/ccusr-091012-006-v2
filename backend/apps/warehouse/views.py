"""
仓库管理视图
"""
import logging
import io
from decimal import Decimal

from django.db.models import Q, Sum
from django.http import HttpResponse
from rest_framework.views import APIView
from rest_framework.permissions import IsAuthenticated
from rest_framework.parsers import MultiPartParser, FormParser, JSONParser
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from apps.core.response import success_response, error_response
from . import services
from .models import (
    Unit, Category, Variety, Goods, StockIn, StockOut, Warning, Approval,
    CustodyPoint, Discrepancy, StockBalance, TransferItem, TransferOrder,
)
from .serializers import (
    UnitSerializer, UnitCreateSerializer,
    CategorySerializer, CategoryCreateSerializer,
    VarietySerializer, VarietyCreateSerializer,
    GoodsSerializer, StockInSerializer, StockOutSerializer,
    WarningSerializer, ApprovalSerializer,
    CustodyPointSerializer, CustodyPointCreateSerializer,
    CancelSerializer, DiscrepancyRaiseSerializer, DiscrepancyResolveSerializer,
    DiscrepancySerializer, DispatchSerializer, ReceiveSerializer,
    SealChangeCreateSerializer, SealChangeSerializer,
    TransferCreateSerializer, TransferEventSerializer,
    TransferOrderDetailSerializer, TransferOrderSerializer,
)

logger = logging.getLogger('apps')


def _first_error_message(errors):
    """从序列化器错误中提取第一条错误信息（兼容嵌套的列表/字典结构）"""
    first = list(errors.values())[0]
    while isinstance(first, (list, dict)):
        first = list(first.values())[0] if isinstance(first, dict) else first[0]
    return str(first)


def _fmt_quantity(value):
    """数量统一格式化为两位小数字符串"""
    return str(Decimal(value).quantize(Decimal('0.01')))


# ==================== 单位管理 ====================

class UnitListView(APIView):
    """单位列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Unit.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        units = queryset[start:end]
        
        serializer = UnitSerializer(units, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建单位"""
        serializer = UnitCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.create(
            name=serializer.validated_data['name'],
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='创建成功')


class UnitDetailView(APIView):
    """单位详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        serializer = UnitCreateSerializer(data=request.data, context={'instance': unit})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit.name = serializer.validated_data['name']
        unit.save()
        
        logger.info(f"User {request.user.username} updated unit {unit.name}")
        
        return success_response(data=UnitSerializer(unit).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除单位"""
        try:
            unit = Unit.objects.get(pk=pk)
        except Unit.DoesNotExist:
            return error_response(message='单位不存在', code=404)
        
        if unit.is_linked:
            return error_response(message='该单位已被关联，无法删除')
        
        name = unit.name
        unit.delete()
        
        logger.info(f"User {request.user.username} deleted unit {name}")
        
        return success_response(message='删除成功')


class UnitBatchDeleteView(APIView):
    """单位批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的单位')
        
        # 只删除未关联的单位
        units = Unit.objects.filter(pk__in=ids)
        deleted_count = 0
        for unit in units:
            if not unit.is_linked:
                unit.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} units")
        
        return success_response(message=f'成功删除 {deleted_count} 个单位')


class UnitAllView(APIView):
    """获取所有单位（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        units = Unit.objects.filter(is_active=True).order_by('name')
        serializer = UnitSerializer(units, many=True)
        return success_response(data=serializer.data)


# ==================== 品类管理 ====================

class CategoryListView(APIView):
    """品类列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Category.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        categories = queryset[start:end]
        
        serializer = CategorySerializer(categories, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品类"""
        serializer = CategoryCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category = Category.objects.create(
            name=serializer.validated_data['name'],
            unit=unit,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='创建成功')


class CategoryDetailView(APIView):
    """品类详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        serializer = CategoryCreateSerializer(data=request.data, context={'instance': category})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0][0]
            return error_response(message=str(first_error))
        
        category.name = serializer.validated_data['name']
        category.unit = Unit.objects.get(pk=serializer.validated_data['unit'])
        category.save()
        
        logger.info(f"User {request.user.username} updated category {category.name}")
        
        return success_response(data=CategorySerializer(category).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品类"""
        try:
            category = Category.objects.get(pk=pk)
        except Category.DoesNotExist:
            return error_response(message='品类不存在', code=404)
        
        if category.is_linked:
            return error_response(message='该品类已被关联，无法删除')
        
        name = category.name
        category.delete()
        
        logger.info(f"User {request.user.username} deleted category {name}")
        
        return success_response(message='删除成功')


class CategoryBatchDeleteView(APIView):
    """品类批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品类')
        
        categories = Category.objects.filter(pk__in=ids)
        deleted_count = 0
        for category in categories:
            if not category.is_linked:
                category.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} categories")
        
        return success_response(message=f'成功删除 {deleted_count} 个品类')


class CategoryAllView(APIView):
    """获取所有品类（用于下拉选择）"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        categories = Category.objects.filter(is_active=True).order_by('name')
        serializer = CategorySerializer(categories, many=True)
        return success_response(data=serializer.data)


# ==================== 品种管理 ====================

class VarietyListView(APIView):
    """品种列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        queryset = Variety.objects.all().order_by('-created_at')
        
        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size
        
        total = queryset.count()
        varieties = queryset[start:end]
        
        serializer = VarietySerializer(varieties, many=True)
        
        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })
    
    def post(self, request):
        """创建品种"""
        serializer = VarietyCreateSerializer(data=request.data)
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        category = Category.objects.get(pk=serializer.validated_data['category'])
        variety = Variety.objects.create(
            name=serializer.validated_data['name'],
            category=category,
            created_by=request.user
        )
        
        logger.info(f"User {request.user.username} created variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='创建成功')


class VarietyDetailView(APIView):
    """品种详情视图"""
    permission_classes = [IsAuthenticated]
    
    def put(self, request, pk):
        """更新品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        serializer = VarietyCreateSerializer(data=request.data, context={'instance': variety})
        if not serializer.is_valid():
            errors = serializer.errors
            first_error = list(errors.values())[0]
            if isinstance(first_error, list):
                first_error = first_error[0]
            return error_response(message=str(first_error))
        
        variety.name = serializer.validated_data['name']
        variety.category = Category.objects.get(pk=serializer.validated_data['category'])
        variety.save()
        
        logger.info(f"User {request.user.username} updated variety {variety.name}")
        
        return success_response(data=VarietySerializer(variety).data, message='更新成功')
    
    def delete(self, request, pk):
        """删除品种"""
        try:
            variety = Variety.objects.get(pk=pk)
        except Variety.DoesNotExist:
            return error_response(message='品种不存在', code=404)
        
        if variety.is_in_stock:
            return error_response(message='该品种已入库，无法删除')
        
        name = variety.name
        variety.delete()
        
        logger.info(f"User {request.user.username} deleted variety {name}")
        
        return success_response(message='删除成功')


class VarietyBatchDeleteView(APIView):
    """品种批量删除视图"""
    permission_classes = [IsAuthenticated]
    
    def post(self, request):
        ids = request.data.get('ids', [])
        if not ids:
            return error_response(message='请选择要删除的品种')
        
        varieties = Variety.objects.filter(pk__in=ids)
        deleted_count = 0
        for variety in varieties:
            if not variety.is_in_stock:
                variety.delete()
                deleted_count += 1
        
        logger.info(f"User {request.user.username} batch deleted {deleted_count} varieties")
        
        return success_response(message=f'成功删除 {deleted_count} 个品种')


class VarietyTemplateView(APIView):
    """品种导入模板下载"""
    permission_classes = []  # 允许匿名访问，通过token参数验证
    
    def get(self, request):
        # 从URL参数获取token进行验证
        from apps.authentication.backends import decode_token
        from apps.authentication.models import User
        
        token = request.query_params.get('token')
        if not token:
            return error_response(message='缺少认证信息', code=401)
        
        payload = decode_token(token)
        if not payload:
            return error_response(message='认证信息无效或已过期', code=401)
        
        try:
            user = User.objects.get(pk=payload['user_id'])
        except User.DoesNotExist:
            return error_response(message='用户不存在', code=401)
        
        wb = Workbook()
        
        # 第一个表格 - 导入模板
        ws1 = wb.active
        ws1.title = '品种导入'
        
        # 设置表头样式
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='4F46E5', end_color='4F46E5', fill_type='solid')
        header_alignment = Alignment(horizontal='center', vertical='center')
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        headers = ['品种', '品类', '单位']
        for col, header in enumerate(headers, 1):
            cell = ws1.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 设置列宽
        ws1.column_dimensions['A'].width = 25
        ws1.column_dimensions['B'].width = 20
        ws1.column_dimensions['C'].width = 15
        
        # 第二个表格 - 品类参考
        ws2 = wb.create_sheet(title='品类参考')
        
        headers2 = ['品类', '单位']
        for col, header in enumerate(headers2, 1):
            cell = ws2.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border
        
        # 填充品类数据
        categories = Category.objects.filter(is_active=True).select_related('unit')
        for row, category in enumerate(categories, 2):
            ws2.cell(row=row, column=1, value=category.name).border = thin_border
            ws2.cell(row=row, column=2, value=category.unit.name).border = thin_border
        
        ws2.column_dimensions['A'].width = 20
        ws2.column_dimensions['B'].width = 15
        
        # 返回Excel文件
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        
        response = HttpResponse(
            output.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename=variety_import_template.xlsx'
        
        return response


class VarietyImportView(APIView):
    """品种导入视图"""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]
    
    def post(self, request):
        if 'file' not in request.FILES:
            return error_response(message='请上传文件')
        
        file = request.FILES['file']
        
        try:
            wb = load_workbook(file)
            ws = wb.active
        except Exception as e:
            return error_response(message='文件格式错误，请上传Excel文件')
        
        # 获取所有品类及其单位
        categories = {c.name: c for c in Category.objects.filter(is_active=True).select_related('unit')}
        
        can_import = []
        cannot_import = []
        
        for row in range(2, ws.max_row + 1):
            variety_name = ws.cell(row=row, column=1).value
            category_name = ws.cell(row=row, column=2).value
            unit_name = ws.cell(row=row, column=3).value
            
            if not variety_name:
                continue
            
            variety_name = str(variety_name).strip()
            category_name = str(category_name).strip() if category_name else ''
            unit_name = str(unit_name).strip() if unit_name else ''
            
            # 验证
            error_msg = None
            
            if not variety_name:
                error_msg = '品种名称不能为空'
            elif len(variety_name) > 20:
                error_msg = '品种名称最多20个字'
            elif not category_name:
                error_msg = '品类不能为空'
            elif category_name not in categories:
                error_msg = f'品类"{category_name}"不存在'
            elif not unit_name:
                error_msg = '单位不能为空'
            elif categories.get(category_name) and categories[category_name].unit.name != unit_name:
                error_msg = f'单位与品类不匹配，应为"{categories[category_name].unit.name}"'
            elif Variety.objects.filter(name=variety_name, category__name=category_name).exists():
                error_msg = '该品种已存在'
            
            if error_msg:
                cannot_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name,
                    'reason': error_msg
                })
            else:
                can_import.append({
                    'row': row,
                    'variety': variety_name,
                    'category': category_name,
                    'unit': unit_name
                })
        
        # 如果是预览请求
        if request.data.get('preview') == 'true':
            return success_response(data={
                'can_import': can_import,
                'cannot_import': cannot_import,
                'can_import_count': len(can_import),
                'cannot_import_count': len(cannot_import)
            })
        
        # 执行导入
        imported_count = 0
        for item in can_import:
            category = categories[item['category']]
            Variety.objects.create(
                name=item['variety'],
                category=category,
                created_by=request.user
            )
            imported_count += 1
        
        logger.info(f"User {request.user.username} imported {imported_count} varieties")
        
        return success_response(
            data={
                'imported_count': imported_count,
                'failed_count': len(cannot_import),
                'failed_items': cannot_import
            },
            message=f'成功导入 {imported_count} 个品种'
        )


# ==================== 其他视图占位 ====================

class DashboardView(APIView):
    """仪表盘视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'message': '仪表盘功能开发中...'
        })


class GoodsListView(APIView):
    """货物列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockInListView(APIView):
    """入库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class StockOutListView(APIView):
    """出库记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class WarningListView(APIView):
    """预警记录列表视图"""
    permission_classes = [IsAuthenticated]
    
    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


class ApprovalListView(APIView):
    """审批记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        return success_response(data={
            'list': [],
            'total': 0,
            'page': 1,
            'page_size': 10
        })


# ==================== 保管点管理 ====================

class CustodyPointListView(APIView):
    """保管点列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = CustodyPoint.objects.all().order_by('-created_at')

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        points = queryset[start:end]

        serializer = CustodyPointSerializer(points, many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """创建保管点"""
        serializer = CustodyPointCreateSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error_message(serializer.errors))

        point = CustodyPoint.objects.create(
            name=serializer.validated_data['name'],
            code=serializer.validated_data['code'],
            address=serializer.validated_data.get('address', ''),
            keeper_name=serializer.validated_data.get('keeper_name', ''),
            created_by=request.user
        )

        logger.info(f"User {request.user.username} created custody point {point.name}")

        return success_response(data=CustodyPointSerializer(point).data, message='创建成功')


class CustodyPointDetailView(APIView):
    """保管点详情视图"""
    permission_classes = [IsAuthenticated]

    def put(self, request, pk):
        """更新保管点"""
        try:
            point = CustodyPoint.objects.get(pk=pk)
        except CustodyPoint.DoesNotExist:
            return error_response(message='保管点不存在', code=404)

        serializer = CustodyPointCreateSerializer(
            data=request.data, context={'instance': point}
        )
        if not serializer.is_valid():
            return error_response(message=_first_error_message(serializer.errors))

        point.name = serializer.validated_data['name']
        point.code = serializer.validated_data['code']
        point.address = serializer.validated_data.get('address', '')
        point.keeper_name = serializer.validated_data.get('keeper_name', '')
        point.save()

        logger.info(f"User {request.user.username} updated custody point {point.name}")

        return success_response(data=CustodyPointSerializer(point).data, message='更新成功')

    def delete(self, request, pk):
        """删除保管点（有库存或调拨往来的不可删除）"""
        try:
            point = CustodyPoint.objects.get(pk=pk)
        except CustodyPoint.DoesNotExist:
            return error_response(message='保管点不存在', code=404)

        if point.origin_transfers.exists() or point.destination_transfers.exists():
            return error_response(message='该保管点存在调拨往来，无法删除')

        nonzero = point.balances.filter(
            Q(available_quantity__gt=0) | Q(reserved_quantity__gt=0)
        ).exists()
        if nonzero:
            return error_response(message='该保管点仍有库存，无法删除')

        name = point.name
        point.balances.all().delete()
        point.delete()

        logger.info(f"User {request.user.username} deleted custody point {name}")

        return success_response(message='删除成功')


class CustodyPointAllView(APIView):
    """获取所有保管点（用于下拉选择）"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        points = CustodyPoint.objects.filter(is_active=True).order_by('name')
        serializer = CustodyPointSerializer(points, many=True)
        return success_response(data=serializer.data)


# ==================== 库存数量查询 ====================

class StockQueryView(APIView):
    """库存数量查询：任意时点区分可用、已预留、在途、待调查数量"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        goods_qs = Goods.objects.filter(is_active=True).select_related(
            'variety__category__unit'
        ).order_by('id')

        goods_id = request.query_params.get('goods_id')
        keyword = request.query_params.get('keyword')
        point_id = request.query_params.get('point_id')
        if goods_id:
            goods_qs = goods_qs.filter(pk=goods_id)
        if keyword:
            goods_qs = goods_qs.filter(
                Q(name__icontains=keyword) | Q(code__icontains=keyword)
            )

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = goods_qs.count()
        goods_page = list(goods_qs[start:end])
        goods_ids = [g.pk for g in goods_page]

        balances = StockBalance.objects.filter(
            goods_id__in=goods_ids
        ).select_related('custody_point')
        open_items = TransferItem.objects.filter(
            goods_id__in=goods_ids, transfer__status='in_transit'
        ).select_related('transfer')
        lost_rows = TransferItem.objects.filter(
            goods_id__in=goods_ids
        ).values('goods_id').annotate(total=Sum('lost_quantity'))
        lost_map = {row['goods_id']: row['total'] for row in lost_rows}

        # 每 (货物, 保管点) 聚合在途与待调查
        balances_by_goods = {}
        for balance in balances:
            balances_by_goods.setdefault(balance.goods_id, []).append(balance)

        transit_out_map = {}
        transit_in_map = {}
        disputed_map = {}
        in_transit_total_map = {}
        disputed_total_map = {}
        extra_point_ids = set()
        for item in open_items:
            goods_pk = item.goods_id
            origin_pk = item.transfer.origin_point_id
            destination_pk = item.transfer.destination_point_id
            in_transit = item.in_transit_quantity
            disputed = item.disputed_quantity
            if in_transit > 0:
                transit_out_map[(goods_pk, origin_pk)] = (
                    transit_out_map.get((goods_pk, origin_pk), Decimal('0')) + in_transit
                )
                transit_in_map[(goods_pk, destination_pk)] = (
                    transit_in_map.get((goods_pk, destination_pk), Decimal('0')) + in_transit
                )
                in_transit_total_map[goods_pk] = (
                    in_transit_total_map.get(goods_pk, Decimal('0')) + in_transit
                )
                extra_point_ids.add(origin_pk)
                extra_point_ids.add(destination_pk)
            if disputed > 0:
                # 待调查数量归起点保管点负责
                disputed_map[(goods_pk, origin_pk)] = (
                    disputed_map.get((goods_pk, origin_pk), Decimal('0')) + disputed
                )
                disputed_total_map[goods_pk] = (
                    disputed_total_map.get(goods_pk, Decimal('0')) + disputed
                )
                extra_point_ids.add(origin_pk)

        point_names = {
            p.pk: p.name
            for p in CustodyPoint.objects.filter(pk__in=extra_point_ids)
        }

        rows = []
        for goods in goods_page:
            goods_balances = balances_by_goods.get(goods.pk, [])
            point_rows = {}
            for balance in goods_balances:
                point_rows[balance.custody_point_id] = {
                    'point_id': balance.custody_point_id,
                    'point_name': balance.custody_point.name,
                    'available': balance.available_quantity,
                    'reserved': balance.reserved_quantity,
                    'in_transit_out': Decimal('0'),
                    'in_transit_in': Decimal('0'),
                    'under_investigation': Decimal('0'),
                }
            # 只有在途/待调查而尚无余额行的保管点（如未签收过的终点）
            for (g_pk, p_pk), qty in list(transit_out_map.items()) + list(transit_in_map.items()) + list(disputed_map.items()):
                if g_pk == goods.pk and p_pk not in point_rows:
                    point_rows[p_pk] = {
                        'point_id': p_pk,
                        'point_name': point_names.get(p_pk, ''),
                        'available': Decimal('0'),
                        'reserved': Decimal('0'),
                        'in_transit_out': Decimal('0'),
                        'in_transit_in': Decimal('0'),
                        'under_investigation': Decimal('0'),
                    }
            for p_pk, row in point_rows.items():
                row['in_transit_out'] = transit_out_map.get((goods.pk, p_pk), Decimal('0'))
                row['in_transit_in'] = transit_in_map.get((goods.pk, p_pk), Decimal('0'))
                row['under_investigation'] = disputed_map.get((goods.pk, p_pk), Decimal('0'))

            point_list = list(point_rows.values())
            if point_id:
                point_list = [r for r in point_list if str(r['point_id']) == str(point_id)]
            for row in point_list:
                for key in ('available', 'reserved', 'in_transit_out',
                            'in_transit_in', 'under_investigation'):
                    row[key] = _fmt_quantity(row[key])

            available_total = sum(
                (b.available_quantity for b in goods_balances), Decimal('0')
            )
            reserved_total = sum(
                (b.reserved_quantity for b in goods_balances), Decimal('0')
            )
            in_transit_total = in_transit_total_map.get(goods.pk, Decimal('0'))
            disputed_total = disputed_total_map.get(goods.pk, Decimal('0'))
            lost_total = lost_map.get(goods.pk, Decimal('0')) or Decimal('0')
            on_hand = available_total + reserved_total + in_transit_total + disputed_total

            rows.append({
                'goods_id': goods.pk,
                'goods_name': goods.name,
                'goods_code': goods.code,
                'unit_name': goods.variety.category.unit.name if goods.variety else '',
                'available': _fmt_quantity(available_total),
                'reserved': _fmt_quantity(reserved_total),
                'in_transit': _fmt_quantity(in_transit_total),
                'under_investigation': _fmt_quantity(disputed_total),
                'lost': _fmt_quantity(lost_total),
                'on_hand': _fmt_quantity(on_hand),
                'points': point_list,
            })

        return success_response(data={
            'list': rows,
            'total': total,
            'page': page,
            'page_size': page_size
        })


# ==================== 调拨管理 ====================

class TransferListView(APIView):
    """调拨单列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = TransferOrder.objects.select_related(
            'origin_point', 'destination_point', 'applicant'
        ).prefetch_related('items__goods').order_by('-created_at')

        status = request.query_params.get('status')
        origin_point = request.query_params.get('origin_point')
        destination_point = request.query_params.get('destination_point')
        keyword = request.query_params.get('keyword')
        if status:
            queryset = queryset.filter(status=status)
        if origin_point:
            queryset = queryset.filter(origin_point_id=origin_point)
        if destination_point:
            queryset = queryset.filter(destination_point_id=destination_point)
        if keyword:
            queryset = queryset.filter(transfer_no__icontains=keyword)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        orders = queryset[start:end]

        serializer = TransferOrderSerializer(orders, many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })

    def post(self, request):
        """调拨申请：起点可用库存转为已预留"""
        serializer = TransferCreateSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error_message(serializer.errors))

        data = serializer.validated_data
        order = services.create_transfer(
            user=request.user,
            origin_point=CustodyPoint.objects.get(pk=data['origin_point']),
            destination_point=CustodyPoint.objects.get(pk=data['destination_point']),
            items=data['items'],
            remark=data.get('remark', ''),
        )

        return success_response(
            data=TransferOrderSerializer(order).data, message='调拨申请已创建'
        )


class TransferDetailView(APIView):
    """调拨单详情视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            order = TransferOrder.objects.select_related(
                'origin_point', 'destination_point', 'applicant', 'dispatched_by'
            ).prefetch_related(
                'items__goods', 'discrepancies', 'seal_changes'
            ).get(pk=pk)
        except TransferOrder.DoesNotExist:
            return error_response(message='调拨单不存在', code=404)

        return success_response(data=TransferOrderDetailSerializer(order).data)


class TransferDispatchView(APIView):
    """起运确认视图：已预留转为在途，记录在途责任人与封签号"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        serializer = DispatchSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error_message(serializer.errors))

        data = serializer.validated_data
        order = services.dispatch_transfer(
            user=request.user,
            order_id=pk,
            lines=data.get('items'),
            carrier_name=data['carrier_name'],
            seal_no=data['seal_no'],
            carrier_phone=data.get('carrier_phone', ''),
            vehicle_no=data.get('vehicle_no', ''),
        )

        return success_response(
            data=TransferOrderSerializer(order).data, message='起运确认成功'
        )


class TransferReceiveView(APIView):
    """到达签收视图：在途转为终点可用，支持部分签收与幂等重试"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        serializer = ReceiveSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error_message(serializer.errors))

        data = serializer.validated_data
        idempotency_key = (
            data.get('idempotency_key')
            or request.headers.get('Idempotency-Key')
            or None
        )
        payload, duplicated = services.receive_transfer(
            user=request.user,
            order_id=pk,
            lines=data.get('items'),
            idempotency_key=idempotency_key,
        )

        message = '重复请求，返回原签收结果' if duplicated else '签收成功'
        return success_response(data=payload, message=message)


class TransferCancelView(APIView):
    """取消调拨视图：释放待起运预留，可将在途数量退回起点"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        serializer = CancelSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error_message(serializer.errors))

        data = serializer.validated_data
        order = services.cancel_transfer(
            user=request.user,
            order_id=pk,
            returns=data.get('returns'),
            reason=data.get('reason', ''),
        )

        return success_response(
            data=TransferOrderSerializer(order).data, message='取消成功'
        )


class TransferSealChangeView(APIView):
    """封签变更视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            order = TransferOrder.objects.get(pk=pk)
        except TransferOrder.DoesNotExist:
            return error_response(message='调拨单不存在', code=404)

        changes = order.seal_changes.all()
        return success_response(data=SealChangeSerializer(changes, many=True).data)

    def post(self, request, pk):
        serializer = SealChangeCreateSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error_message(serializer.errors))

        change = services.change_seal(
            user=request.user,
            order_id=pk,
            new_seal_no=serializer.validated_data['new_seal_no'],
            reason=serializer.validated_data.get('reason', ''),
        )

        return success_response(
            data=SealChangeSerializer(change).data, message='封签变更成功'
        )


class TransferDiscrepancyView(APIView):
    """调拨单差异视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            order = TransferOrder.objects.get(pk=pk)
        except TransferOrder.DoesNotExist:
            return error_response(message='调拨单不存在', code=404)

        discrepancies = order.discrepancies.all()
        return success_response(data=DiscrepancySerializer(discrepancies, many=True).data)

    def post(self, request, pk):
        """差异登记：在途数量转为待调查"""
        serializer = DiscrepancyRaiseSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error_message(serializer.errors))

        data = serializer.validated_data
        discrepancy = services.raise_discrepancy(
            user=request.user,
            order_id=pk,
            item_id=data['item_id'],
            quantity=data['quantity'],
            reason=data['reason'],
        )

        return success_response(
            data=DiscrepancySerializer(discrepancy).data, message='差异已登记'
        )


class TransferEventListView(APIView):
    """调拨事件台账视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request, pk):
        try:
            order = TransferOrder.objects.get(pk=pk)
        except TransferOrder.DoesNotExist:
            return error_response(message='调拨单不存在', code=404)

        events = order.events.all()
        return success_response(data=TransferEventSerializer(events, many=True).data)


class DiscrepancyListView(APIView):
    """差异记录列表视图"""
    permission_classes = [IsAuthenticated]

    def get(self, request):
        queryset = Discrepancy.objects.select_related(
            'transfer', 'item__goods', 'raised_by', 'resolved_by'
        ).order_by('-created_at')

        status = request.query_params.get('status')
        transfer = request.query_params.get('transfer')
        if status:
            queryset = queryset.filter(status=status)
        if transfer:
            queryset = queryset.filter(transfer_id=transfer)

        page = int(request.query_params.get('page', 1))
        page_size = int(request.query_params.get('page_size', 10))
        start = (page - 1) * page_size
        end = start + page_size

        total = queryset.count()
        discrepancies = queryset[start:end]

        serializer = DiscrepancySerializer(discrepancies, many=True)

        return success_response(data={
            'list': serializer.data,
            'total': total,
            'page': page,
            'page_size': page_size
        })


class DiscrepancyResolveView(APIView):
    """差异处理视图：待调查数量补交终点、退回起点或核销损失"""
    permission_classes = [IsAuthenticated]

    def post(self, request, pk):
        serializer = DiscrepancyResolveSerializer(data=request.data)
        if not serializer.is_valid():
            return error_response(message=_first_error_message(serializer.errors))

        data = serializer.validated_data
        discrepancy = services.resolve_discrepancy(
            user=request.user,
            discrepancy_id=pk,
            delivered=data['delivered_quantity'],
            returned=data['returned_quantity'],
            lost=data['lost_quantity'],
            remark=data.get('remark', ''),
        )

        return success_response(
            data=DiscrepancySerializer(discrepancy).data, message='差异处理成功'
        )
